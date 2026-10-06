"""Construye un CSV de cosecha de paltas a partir de planillas semanales.

Uso:
    python -m ia.preparar_dataset "C:\\ruta\\a\\carpeta\\2026"

Por seguridad, el script solo toma archivos cuyo nombre contiene "Cosecha Paltas"
(plural). De esta forma no mezcla automáticamente Anillado, limón u otras labores.
"""

from __future__ import annotations

import re
import sys
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

DIAS = {
    "lunes": 1,
    "martes": 2,
    "miércoles": 3,
    "miercoles": 3,
    "jueves": 4,
    "viernes": 5,
    "sábado": 6,
    "sabado": 6,
    "domingo": 7,
}


def limpiar_nombre(valor) -> str | None:
    if valor is None:
        return None
    return re.sub(r"\s+", " ", str(valor)).strip() or None


def fechas_esperadas(hoja, semana: int, anio: int = 2026) -> dict[int, date]:
    etiquetas = [
        str(hoja.cell(1, columna).value or "").strip().lower()
        for columna in range(5, 11)
    ]

    posicion_lunes = next(
        (indice for indice, texto in enumerate(etiquetas) if texto == "lunes"),
        None,
    )

    fechas = {}
    for indice, columna in enumerate(range(5, 11)):
        etiqueta = etiquetas[indice]
        if etiqueta not in DIAS:
            valor = hoja.cell(2, columna).value
            if isinstance(valor, datetime):
                fechas[columna] = valor.date()
            continue

        semana_iso = semana
        if posicion_lunes is not None and indice < posicion_lunes and DIAS[etiqueta] >= 5:
            semana_iso = semana - 1

        fechas[columna] = date.fromisocalendar(anio, semana_iso, DIAS[etiqueta])

    return fechas


def tarifa_por_fecha(semana: int, fecha: date) -> float | None:
    # Tarifas observadas en las planillas históricas entregadas.
    if semana == 31:
        return 55.0
    if semana == 32:
        return 60.0
    if semana == 33:
        return 60.0 if fecha.weekday() in (0, 1) else 70.0
    if semana in (34, 35, 36):
        return 75.0
    if semana == 37:
        return 75.0 if fecha <= date(2026, 9, 7) else 80.0
    if semana >= 38:
        return 80.0
    return None


def leer_archivo(ruta: Path) -> list[dict]:
    coincidencia = re.search(r"Semana\s+(\d+)", ruta.name, re.IGNORECASE)
    if not coincidencia:
        return []

    semana = int(coincidencia.group(1))
    libro = load_workbook(ruta, data_only=False)
    nombre_hoja = "Cosecha" if "Cosecha" in libro.sheetnames else libro.sheetnames[0]
    hoja = libro[nombre_hoja]
    fechas = fechas_esperadas(hoja, semana)

    registros = []
    for fila in range(3, 80):
        nombre = limpiar_nombre(hoja.cell(fila, 2).value)
        valores = [hoja.cell(fila, columna).value for columna in fechas]
        tiene_numero = any(
            isinstance(valor, (int, float)) and not isinstance(valor, bool)
            for valor in valores
        )

        # La primera fila completamente vacía marca el fin de la lista de cosechadores.
        if not nombre and not tiene_numero:
            break
        if not nombre or not tiene_numero:
            continue

        for columna, fecha in fechas.items():
            kilos = hoja.cell(fila, columna).value
            if not isinstance(kilos, (int, float)) or isinstance(kilos, bool):
                continue
            if kilos <= 0 or kilos > 5000:
                continue

            original = hoja.cell(2, columna).value
            fecha_original = original.date() if isinstance(original, datetime) else None

            registros.append(
                {
                    "archivo": ruta.name,
                    "semana": semana,
                    "fecha": fecha.isoformat(),
                    "fecha_original": fecha_original.isoformat() if fecha_original else "",
                    "fecha_corregida": bool(fecha_original and fecha_original != fecha),
                    "trabajador": nombre,
                    "kilos": round(float(kilos), 2),
                    "tarifa_kg": tarifa_por_fecha(semana, fecha),
                }
            )

    return registros


def construir(carpeta: Path, salida: Path) -> pd.DataFrame:
    archivos = sorted(
        ruta
        for ruta in carpeta.rglob("*.xlsx")
        if "cosecha paltas" in ruta.name.lower()
    )

    registros = []
    for archivo in archivos:
        registros.extend(leer_archivo(archivo))

    datos = pd.DataFrame(registros)
    if datos.empty:
        raise RuntimeError("No se encontraron registros de Cosecha Paltas.")

    datos = datos.sort_values(["fecha", "trabajador"]).reset_index(drop=True)
    salida.parent.mkdir(parents=True, exist_ok=True)
    datos.to_csv(salida, index=False, encoding="utf-8-sig")
    return datos


def main() -> None:
    if len(sys.argv) < 2:
        print("Debes indicar la carpeta que contiene las planillas.")
        print('Ejemplo: python -m ia.preparar_dataset "C:\\Datos\\2026"')
        raise SystemExit(1)

    carpeta = Path(sys.argv[1]).expanduser().resolve()
    salida = Path(__file__).resolve().parent / "dataset_cosecha_paltas_2026.csv"
    datos = construir(carpeta, salida)
    print(f"Dataset creado: {salida}")
    print(f"Registros: {len(datos)}")
    print(f"Trabajadores: {datos['trabajador'].nunique()}")
    print(f"Fechas corregidas: {int(datos['fecha_corregida'].sum())}")


if __name__ == "__main__":
    main()
